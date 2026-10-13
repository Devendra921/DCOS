"""
Extract a DCO inventory workbook (template / backup layout) into a compact JSON dataset that the
app loads at start-up (src/mock/datasets/neocloud/<estate>.json). Mapping into DCOS records happens
in TypeScript (src/mock/datasets/dcoImport.ts) so the same rules can later serve an upload screen.

  python3 scripts/import/extract_dco_workbook.py <workbook.xlsx> <estate-code> <currency> [--all-ports]

Per sheet: {"cols": [...], "rows": [[...], ...]}. Calculated columns (header "… (calc)") are dropped,
dates become YYYY-MM-DD, blanks become null. Unless --all-ports, Interfaces keeps only ports that are
up, reserved or cabled (spare admin-down ports are left out, as the template allows).
"""
import json, sys, datetime
from openpyxl import load_workbook

SHEETS = ['Sites', 'Floors', 'Rooms', 'Racks', 'Power_Feeds', 'UPS', 'Generators', 'DC_Plant', 'Cooling_Units', 'CDUs',
          'Sensors', 'Devices', 'Modules', 'Interfaces', 'Links', 'GPU_Clusters', 'GPU_Nodes', 'Customers', 'Services',
          'Collectors', 'Snapshot_Alerts', 'Snapshot_Incidents', 'Snapshot_Changes']

def cell(v):
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d') if v.time() == datetime.time() else v.strftime('%Y-%m-%dT%H:%M:%SZ')
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v

def main():
    path, estate, currency = sys.argv[1], sys.argv[2], sys.argv[3]
    all_ports = '--all-ports' in sys.argv
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {'estate': estate, 'currency': currency, 'source': path.split('/')[-1], 'sheets': {}}
    for name in SHEETS:
        if name not in wb.sheetnames: continue
        it = wb[name].iter_rows(values_only=True)
        hdr = [(h or '').replace(' *', '').strip() for h in next(it)]
        keep = [i for i, h in enumerate(hdr) if h and '(calc)' not in h]
        rows = [[cell(r[i]) for i in keep] for r in it if any(v not in (None, '') for v in r)]
        out['sheets'][name] = {'cols': [hdr[i] for i in keep], 'rows': rows}
    if not all_ports:
        links = out['sheets']['Links']; lc = links['cols']
        cabled = set()
        for r in links['rows']:
            cabled.add((r[lc.index('a_device_id')], r[lc.index('a_if_name')])); cabled.add((r[lc.index('z_device_id')], r[lc.index('z_if_name')]))
        ifs = out['sheets']['Interfaces']; ic = ifs['cols']
        di, ni, si = ic.index('device_id'), ic.index('if_name'), ic.index('admin_state')
        before = len(ifs['rows'])
        ifs['rows'] = [r for r in ifs['rows'] if r[si] != 'admin-down' or (r[di], r[ni]) in cabled]
        out['portsDropped'] = before - len(ifs['rows'])
    dest = f"src/mock/datasets/neocloud/{estate}.json"
    with open(dest, 'w') as f: json.dump(out, f, separators=(',', ':'))
    print(dest, {k: len(v['rows']) for k, v in out['sheets'].items()}, 'ports dropped', out.get('portsDropped'))

if __name__ == '__main__': main()
